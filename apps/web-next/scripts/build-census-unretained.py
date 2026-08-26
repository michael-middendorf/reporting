#!/usr/bin/env python3
"""Rebuild unretained.json: last census raw − intervening graduates − this-term registration.

Student-level output is gitignored. Registration is the all-student count for the current term.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

CENSUS_DIR = Path("/Users/regina.m.pirolli/Documents/PERSONAL Census Files")
PRIOR = CENSUS_DIR / "Fall 2025 Census Final 9.23.25.xlsx"
REGISTRATION = Path("/Users/regina.m.pirolli/Downloads/Registration Status-23.xlsx")
GRADUATES = CENSUS_DIR / "FA26 Census" / "FA26_graduates_2025_2026.xlsx"
OUT = Path(__file__).resolve().parents[1] / "src/lib/reports/census/data/unretained.json"

NON_DS = {"DUAL", "EACO", "ND"}


def norm(v) -> str:
    return str(v or "").strip().upper()


def is_ay_grad(year, term) -> bool:
    y = str(year or "").strip()
    t = str(term or "").strip().upper()
    return (y == "2025" and t == "FALL") or (y == "2026" and t in {"SPRING", "SUMMER"})


def header_map(row) -> dict[str, int]:
    out = {}
    for i, v in enumerate(row):
        key = str(v or "").strip()
        if key and key not in out:
            out[key] = i
    return out


def dash(v) -> str:
    t = str(v).strip() if v is not None else ""
    return t if t else "-"


def load_registration_ids(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    ids: set[str] = set()
    header = None
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and str(row[0] or "").strip() == "ID":
                header = header_map(row)
            continue
        pcid = norm(row[0])
        if pcid.startswith("P"):
            ids.add(pcid)
    wb.close()
    return ids


def load_graduate_ids(path: Path) -> set[str]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = header_map(next(rows))
    ids: set[str] = set()
    for row in rows:
        pcid = norm(row[header["PCID"]])
        if not pcid:
            continue
        if is_ay_grad(row[header["Graduation Year"]], row[header["Graduation Term"]]):
            ids.add(pcid)
    wb.close()
    return ids


def load_prior(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["FA25 Census Raw Data"]
    rows = ws.iter_rows(values_only=True)
    header = header_map(next(rows))
    out = []
    for row in rows:
        pcid = norm(row[header["PCID"]])
        if not pcid.startswith("P"):
            continue
        sport = dash(row[header.get("Athletic Assoc. (Athlet.)", header.get("Athletic Assoc.", 18))])
        out.append(
            {
                "pcid": row[header["PCID"]],
                "first": str(row[header["First_Name"]] or "").strip(),
                "last": str(row[header["Last_Name"]] or "").strip(),
                "retention": "Not Retained",
                "gpa": row[header["Cumulative GPA"]] if row[header["Cumulative GPA"]] not in (None, "") else "-",
                "fsp": dash(row[header["FSP"]] if "FSP" in header else ""),
                "household": dash(row[header["Household"]] if "Household" in header else ""),
                "marines": dash(row[header["Mary's Marines"]] if "Mary's Marines" in header else ""),
                "honors": dash(row[header["Honors"]] if "Honors" in header else ""),
                "sport": sport,
                "state": dash(row[header["State"]]),
                "gender": dash(row[header["Gender"]]),
                "classLevel": dash(row[header["Class_Level"]]),
                "collegeAttend": dash(row[header["College_Attend"]]),
                "incomingReturning": dash(row[header["Incoming/Returning"]]),
                "program": dash(row[header["Program"]]),
            }
        )
    wb.close()
    return out


def main() -> None:
    registered = load_registration_ids(REGISTRATION)
    graduates = load_graduate_ids(GRADUATES)
    prior = load_prior(PRIOR)
    unretained = [
        s
        for s in prior
        if norm(s["pcid"]) not in registered and norm(s["pcid"]) not in graduates
    ]
    ug = [
        s
        for s in unretained
        if s.get("program") == "UNDERG" and s.get("collegeAttend") not in NON_DS
    ]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated": date.today().isoformat(),
        "term": "Fall 2026",
        "prior": "Fall 2025",
        "definition": "FA25 census raw minus 2025-26 Fall/Spring/Summer graduates minus FA26 registration.",
        "count": len(unretained),
        "ugDegreeSeeking": len(ug),
        "registered": len(registered),
        "graduates": len(graduates),
        "priorCensus": len(prior),
        "students": unretained,
    }
    OUT.write_text(json.dumps(payload, default=str))
    print(
        f"Wrote {OUT}\n"
        f"  prior census {len(prior)}\n"
        f"  registered this term {len(registered)}\n"
        f"  intervening graduates {len(graduates)}\n"
        f"  unretained {len(unretained)} (UG degree-seeking {len(ug)})"
    )


if __name__ == "__main__":
    main()
