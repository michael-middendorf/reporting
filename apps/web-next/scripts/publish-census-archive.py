#!/usr/bin/env python3
"""Publish the local Census folder zip to the reporting data bucket.

Reads a downloaded zip only. Does not open or change Drive files.
Workbook bytes and sheet previews stay out of git.

  python3 scripts/publish-census-archive.py --zip "$HOME/Downloads/Census (1).zip" --out /tmp/census-archive
  python3 scripts/publish-census-archive.py --zip "$HOME/Downloads/Census (1).zip" --out /tmp/census-archive --upload
"""

from __future__ import annotations

import argparse
import json
import math
import re
import subprocess
import zipfile
from datetime import date, datetime
from pathlib import Path

MAX_ROWS = 4000
MAX_COLS = 120
EMPTY_RUN = 40

CONTENT_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xls": "application/vnd.ms-excel",
    ".xlk": "application/vnd.ms-excel",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".pbix": "application/octet-stream",
    ".msg": "application/vnd.ms-outlook",
}

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_encrypted_workbook(data: bytes) -> bool:
    for marker in (b"EncryptedPackage", b"DRMEncrypted"):
        if marker in data or marker.decode().encode("utf-16-le") in data:
            return True
    return False


def slug(text: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return cleaned[:160] or "file"


def term_of(folder: str) -> tuple[str, int | None]:
    match = re.match(r"^(FA|SP)\s*(\d{2})\b", folder.strip(), re.I)
    if not match:
        return "other", None
    year = 2000 + int(match.group(2))
    season = "fall" if match.group(1).upper() == "FA" else "spring"
    return season, year


def is_blank(text: str) -> bool:
    return text.strip() == ""


def number_text(value: float) -> str:
    if not math.isfinite(value):
        return ""
    if abs(value - round(value)) < 1e-9 and abs(value) < 1e15:
        return str(int(round(value)))
    return f"{value:.4f}".rstrip("0").rstrip(".")


def percent_text(value: float, fmt: str) -> str:
    pct = value * 100
    if "." in fmt.split("%", 1)[0]:
        return f"{pct:.1f}%"
    return f"{pct:.0f}%"


def text_of(value, fmt: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and fmt and "%" in fmt:
        return percent_text(value, fmt)
    if isinstance(value, float):
        return number_text(value)
    if isinstance(value, int):
        return str(value)
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text[:2000]


def trim_rows(rows: list[list[str]], truncated: bool) -> tuple[list[list[str]], bool]:
    while rows and all(is_blank(cell) for cell in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for index, cell in enumerate(row):
            if not is_blank(cell):
                width = max(width, index + 1)
    trimmed = [row[:width] + [""] * (width - len(row[:width])) for row in rows]
    return trimmed, truncated


def limit_rows(raw_rows) -> tuple[list[list[str]], bool]:
    kept: list[list[str]] = []
    empty_run = 0
    truncated = False
    started = False
    scanned = 0
    for values in raw_rows:
        scanned += 1
        # Styled workbooks sometimes claim a million empty rows.
        if scanned > 12000:
            truncated = True
            break
        row = values[:MAX_COLS]
        blank = all(is_blank(cell) for cell in row)
        if blank:
            empty_run += 1
            if not started and empty_run >= 80:
                break
            if started and empty_run >= EMPTY_RUN:
                break
            if started:
                kept.append(row)
            continue
        started = True
        empty_run = 0
        kept.append(row)
        if len(kept) >= MAX_ROWS:
            truncated = True
            break
    return trim_rows(kept, truncated)


def preview_xlsx(data: bytes) -> list[tuple[str, dict]]:
    import io
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    try:
        for name in book.sheetnames:
            ws = book[name]

            def rows():
                for row in ws.iter_rows(max_col=MAX_COLS):
                    yield [text_of(cell.value, getattr(cell, "number_format", "") or "") for cell in row]

            body, truncated = limit_rows(rows())
            width = max((len(row) for row in body), default=0)
            sheets.append((name, {
                "name": name,
                "rows": body,
                "rowCount": len(body),
                "colCount": width,
                "truncated": truncated,
            }))
    finally:
        book.close()
    return sheets


def preview_xls(data: bytes) -> list[tuple[str, dict]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets = []
    for ws in book.sheets():
        def rows(sheet=ws):
            for r in range(sheet.nrows):
                out = []
                for c in range(min(sheet.ncols, MAX_COLS)):
                    cell = sheet.cell(r, c)
                    value = cell.value
                    fmt = ""
                    try:
                        xf = book.xf_list[cell.xf_index]
                        fmt = book.format_map[xf.format_key].format_str or ""
                    except Exception:
                        fmt = ""
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        except Exception:
                            value = ""
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = ""
                    out.append(text_of(value, fmt))
                yield out

        body, truncated = limit_rows(rows())
        width = max((len(row) for row in body), default=0)
        sheets.append((ws.name, {
            "name": ws.name,
            "rows": body,
            "rowCount": len(body),
            "colCount": width,
            "truncated": truncated,
        }))
    return sheets


def kind_for(ext: str, sheets: list) -> str:
    if sheets:
        return "workbook"
    if ext == ".pdf":
        return "pdf"
    return "other"


def iter_zip(path: Path):
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            name = info.filename.replace("\\", "/")
            if name.startswith("__MACOSX/") or "/._" in f"/{name}" or name.endswith(".DS_Store"):
                continue
            base = name.rsplit("/", 1)[-1]
            if not base or base.startswith("~$") or base.startswith("."):
                continue
            rel = name
            if rel.lower().startswith("census/"):
                rel = rel.split("/", 1)[1]
            yield rel, archive.read(info)


def build(zip_path: Path, out: Path) -> dict:
    files_dir = out / "files"
    sheets_dir = out / "sheets"
    files_dir.mkdir(parents=True, exist_ok=True)
    used: set[str] = set()
    files = []
    for rel, data in iter_zip(zip_path):
        folder = str(Path(rel).parent)
        if folder == ".":
            folder = ""
        name = Path(rel).name
        ext = Path(name).suffix.lower()
        file_id = slug(rel)
        if file_id in used:
            file_id = slug(f"{rel}-{len(used)}")
        used.add(file_id)
        (files_dir / file_id).write_bytes(data)
        sheets = []
        note = ""
        content_type = CONTENT_TYPES.get(ext, "application/octet-stream")
        try:
            ole = data.startswith(OLE_MAGIC)
            if ext == ".xlsx" and ole:
                sheets = preview_xls(data)
            elif ext in {".xlsx", ".xlsm"} or (ext == ".xlk" and data.startswith(b"PK")):
                sheets = preview_xlsx(data)
            elif ext in {".xls", ".xlk"}:
                sheets = preview_xls(data)
        except Exception as exc:
            if is_encrypted_workbook(data):
                note = "This workbook is encrypted. Download it and open it in Excel."
            else:
                note = "This workbook could not be opened here. Download the file."
            sheets = []
            print(f"  preview failed {rel}: {type(exc).__name__}: {exc}", flush=True)
        season, year = term_of(folder)
        entry = {
            "id": file_id,
            "folder": folder,
            "name": name,
            "bytes": len(data),
            "ext": ext.lstrip("."),
            "kind": kind_for(ext, sheets),
            "season": season,
            "year": year,
            "contentType": content_type,
            "sheets": [sheet_name for sheet_name, _sheet in sheets],
        }
        if note:
            entry["note"] = note
        sheet_out = sheets_dir / file_id
        if sheets:
            sheet_out.mkdir(parents=True, exist_ok=True)
            for index, (_sheet_name, payload) in enumerate(sheets):
                (sheet_out / f"{index}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        files.append(entry)
        print(f"{entry['kind']:8} {rel} sheets={len(sheets)}", flush=True)
    manifest = {
        "generated": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source": "Local Census folder download. Drive files were not changed.",
        "files": files,
    }
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def upload(out: Path, bucket: str, manifest: dict) -> None:
    prefix = f"s3://{bucket}/cache/census/archive"
    jobs = []
    for entry in manifest["files"]:
        jobs.append((out / "files" / entry["id"], f"{prefix}/files/{entry['id']}", entry["contentType"]))
        for index in range(len(entry["sheets"])):
            jobs.append((
                out / "sheets" / entry["id"] / f"{index}.json",
                f"{prefix}/sheets/{entry['id']}/{index}.json",
                "application/json; charset=utf-8",
            ))

    def copy(job):
        src, dest, content_type = job
        subprocess.check_call([
            "aws", "s3", "cp", str(src), dest,
            "--content-type", content_type,
            "--cache-control", "private, max-age=300",
            "--only-show-errors",
        ])

    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(copy, jobs))
    copy((out / "manifest.json", f"{prefix}/manifest.json", "application/json; charset=utf-8"))
    print(f"Published {len(jobs) + 1} objects to {prefix}/", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--upload", action="store_true")
    parser.add_argument("--bucket", default="amu-reporting-production-reportdatabucket-rncvucke")
    args = parser.parse_args()
    manifest = build(args.zip, args.out)
    print(f"Wrote {len(manifest['files'])} files to {args.out}", flush=True)
    if args.upload:
        upload(args.out, args.bucket, manifest)


if __name__ == "__main__":
    main()
