#!/usr/bin/env python3
"""Join admissions / withdrawal fields onto FA25 not-retained PCIDs.

Writes gitignored data/unretained-enrichment.json. Rebuild:
  python3 scripts/build-unretained-enrichment.py
"""
from __future__ import annotations

import csv
import json
import math
import re
import urllib.request
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

DOWNLOADS = Path("/Users/regina.m.pirolli/Downloads")
DATA = Path(__file__).resolve().parents[1] / "src/lib/reports/census/data"
OUT = DATA / "unretained-enrichment.json"

ML_TRAINING = DOWNLOADS / "ML Training Data -_ Qlik 20260709-120623.xlsx"
DEPOSITS = DOWNLOADS / "FA25 Deposits With Scholarship Info 20260410-105249.xlsx"
HS_NAME = DOWNLOADS / "Qlik Sense - Untitled - Table - September 18, 2025.xlsx"
WITHDRAWAL = next(
    (
        p
        for p in (
            DOWNLOADS / "Institutional Withdrawal Exit Interview 2024-2026 (2).xlsx",
            DOWNLOADS / "Institutional Withdrawal Exit Interview 2024-2026 (1).xlsx",
            DOWNLOADS / "Institutional Withdrawal Exit Interview 2024-2026.xlsx",
        )
        if p.exists() and not p.name.startswith("~$")
    ),
    DOWNLOADS / "Institutional Withdrawal Exit Interview 2024-2026 (2).xlsx",
)
ADM_ML = DOWNLOADS / "FA25 Admissions Report ML -_ Qlik 20260408-160056.xlsx"
INCOMING = DOWNLOADS / "FA25 Incoming Students Census 20260908-102528.xlsx"
GOLD_FTCA = Path("/Users/regina.m.pirolli/Documents/Census Files/Fall 2025 Census Final 9.23.25.xlsx")
GOLD_FTCA_SHEET = "F24-F25 FTCA Retention Raw Data"
PRIOR_JSON = DATA / "prior.json"
GEO_CACHE = DATA / "geo"
GAZ_URL = "https://www2.census.gov/geo/docs/maps-data/data/gazetteer/2024_Gazetteer/2024_Gaz_zcta_national.zip"
CITIES_URL = "https://raw.githubusercontent.com/kelvins/US-Cities-Database/main/csv/us_cities.csv"
# Ave Maria University, 5050 Ave Maria Blvd.
CAMPUS_LAT = 26.3365
CAMPUS_LON = -81.4367

MONEY_RE = re.compile(r"[^\d.\-]")


def norm_id(v) -> str:
    t = str(v or "").strip().upper()
    if not t or t in {"-", "NONE", "NULL"}:
        return ""
    if t.isdigit():
        return "P" + t.zfill(9)
    return t


def money(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = MONEY_RE.sub("", str(v).replace(",", ""))
    if t in {"", ".", "-"}:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def blank(v) -> str:
    if v is None:
        return ""
    t = str(v).strip()
    return "" if t in {"-", "None", "null"} else t


def load_sheet(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(h or "").strip() for h in rows[0]]
    return headers, rows[1:]


def col(headers, *names):
    lower = {h.lower(): i for i, h in enumerate(headers)}
    for n in names:
        if n.lower() in lower:
            return lower[n.lower()]
    return None


def put(dst: dict, key: str, val, overwrite: bool = False):
    if val is None or val == "":
        return
    if overwrite or dst.get(key) in (None, ""):
        dst[key] = val


def score(v, sat: bool = False):
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return blank(v) or None
    if n == 0:
        return None
    if sat and n < 400:
        return None
    return int(n) if n == int(n) else n


def haversine_miles(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 3958.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlmb = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return r * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def zip5(v) -> str:
    t = str(v or "").strip()
    if not t or t.lower().startswith("no zi"):
        return ""
    digits = re.sub(r"\D", "", t)
    return digits[:5] if len(digits) >= 5 else ""


CITY_ALIASES = {
    "virginia bch": "virginia beach",
    "salt lake cty": "salt lake city",
    "pembroke pnes": "pembroke pines",
    "commerce twp": "commerce township",
    "ft myers": "fort myers",
    "ft lauderdale": "fort lauderdale",
    "st louis": "st louis",
    "saint louis": "st louis",
    "saint johns": "st johns",
}


def norm_city(v: str) -> str:
    t = re.sub(r"[.']", "", (v or "").lower())
    t = re.sub(r"\s+", " ", t).strip()
    t = t.replace("saint ", "st ")
    return CITY_ALIASES.get(t, t)


def ensure_geo() -> tuple[dict[str, tuple[float, float]], dict[tuple[str, str], tuple[float, float]]]:
    GEO_CACHE.mkdir(parents=True, exist_ok=True)
    gaz_path = GEO_CACHE / "2024_Gaz_zcta_national.txt"
    if not gaz_path.exists():
        raw = urllib.request.urlopen(GAZ_URL, timeout=60).read()
        with zipfile.ZipFile(BytesIO(raw)) as zf:
            name = next(n for n in zf.namelist() if n.endswith(".txt"))
            gaz_path.write_bytes(zf.read(name))
    zips: dict[str, tuple[float, float]] = {}
    with gaz_path.open(encoding="utf-8", errors="replace") as f:
        next(f)
        for line in f:
            parts = line.split()
            if len(parts) < 3:
                continue
            zips[parts[0].zfill(5)] = (float(parts[-2]), float(parts[-1]))
    cities_path = GEO_CACHE / "us_cities.csv"
    if not cities_path.exists():
        cities_path.write_bytes(urllib.request.urlopen(CITIES_URL, timeout=60).read())
    cities: dict[tuple[str, str], tuple[float, float]] = {}
    with cities_path.open(encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            city = norm_city(row.get("CITY") or "")
            st = (row.get("STATE_CODE") or "").strip().upper()
            try:
                lat, lon = float(row["LATITUDE"]), float(row["LONGITUDE"])
            except (KeyError, TypeError, ValueError):
                continue
            if city and st and (city, st) not in cities:
                cities[(city, st)] = (lat, lon)
    cities.setdefault(("miramar", "FL"), (25.986, -80.303))
    cities.setdefault(("commerce township", "MI"), (42.569, -83.488))
    return zips, cities


def miles_from_campus(
    zips: dict[str, tuple[float, float]],
    cities: dict[tuple[str, str], tuple[float, float]],
    zip_code: str = "",
    city: str = "",
    state: str = "",
):
    city_n = norm_city(city)
    st = (state or "").strip().upper()
    if city_n == "ave maria" or (city_n.startswith("ave maria") and st in {"FL", ""}):
        return 0
    if zip_code and zip_code in zips:
        lat, lon = zips[zip_code]
        return int(round(haversine_miles(CAMPUS_LAT, CAMPUS_LON, lat, lon)))
    if city_n and st and (city_n, st) in cities:
        lat, lon = cities[(city_n, st)]
        return int(round(haversine_miles(CAMPUS_LAT, CAMPUS_LON, lat, lon)))
    return None


def term_rank(term: str) -> int:
    t = (term or "").lower()
    if "2025" in t and "fall" in t:
        return 3
    if "2024" in t and "fall" in t:
        return 2
    if "2026" in t and "fall" in t:
        return 1
    return 0


def main() -> None:
    by_id: dict[str, dict] = {}

    if ML_TRAINING.exists():
        headers, rows = load_sheet(ML_TRAINING)
        i_id = col(headers, "PCID")
        i_term = col(headers, "Application Term")
        i_total = col(headers, "Interview Total Rating")
        i_faith = col(headers, "Interview - Faith (24/25)", "Interview - Faith Rating (23/24)")
        i_type = col(headers, "Ave type")
        scored = []
        for r in rows:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            scored.append((term_rank(str(r[i_term] or "") if i_term is not None else ""), pid, r))
        scored.sort(reverse=True)
        seen = set()
        for _, pid, r in scored:
            if pid in seen:
                continue
            seen.add(pid)
            rec = by_id.setdefault(pid, {"pcid": pid})
            put(rec, "interviewTotal", r[i_total] if i_total is not None else None)
            put(rec, "interviewFaith", r[i_faith] if i_faith is not None else None)
            put(rec, "schoolType", blank(r[i_type] if i_type is not None else ""))

    if GOLD_FTCA.exists():
        wb_gold = load_workbook(GOLD_FTCA, data_only=True, read_only=True)
        ws_gold = wb_gold[GOLD_FTCA_SHEET] if GOLD_FTCA_SHEET in wb_gold.sheetnames else wb_gold[wb_gold.sheetnames[0]]
        gold_rows = list(ws_gold.iter_rows(values_only=True))
        wb_gold.close()
        headers = [str(h or "").strip() for h in gold_rows[0]]
        i_id = col(headers, "PCID")
        i_gpa = col(headers, "HS GPA")
        for r in gold_rows[1:]:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            rec = by_id.setdefault(pid, {"pcid": pid})
            put(rec, "hsGpa", score(r[i_gpa] if i_gpa is not None else None))

    if ADM_ML.exists():
        headers, rows = load_sheet(ADM_ML)
        i_id = col(headers, "Person PowerCampus ID", "PCID")
        i_major = col(headers, "Major")
        i_ath = col(headers, "Athletic Scholarship Recieved")
        i_unf = col(headers, "Unfunded Aid")
        i_tot = col(headers, "Person Interview - Total Rating")
        for r in rows:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            rec = by_id.setdefault(pid, {"pcid": pid})
            put(rec, "applicationMajor", blank(r[i_major] if i_major is not None else ""))
            put(rec, "athleticAid", money(r[i_ath] if i_ath is not None else None))
            put(rec, "unfundedAid", money(r[i_unf] if i_unf is not None else None))
            put(rec, "interviewTotal", r[i_tot] if i_tot is not None else None)

    if DEPOSITS.exists():
        headers, rows = load_sheet(DEPOSITS)
        i_id = col(headers, "PCID")
        i_ath = col(headers, "Athletic")
        i_tot = col(headers, "Total Scholarship Recieved")
        i_acad = col(headers, "Academic")
        i_hon = col(headers, "Honors")
        i_fsp = col(headers, "FSP")
        i_inst = col(headers, "Institute")
        i_oth = col(headers, "Other Scholarships")
        for r in rows:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            rec = by_id.setdefault(pid, {"pcid": pid})
            ath = money(r[i_ath] if i_ath is not None else None)
            tot = money(r[i_tot] if i_tot is not None else None)
            other_parts = [
                money(r[c] if c is not None else None)
                for c in (i_acad, i_hon, i_fsp, i_inst, i_oth)
            ]
            other = sum(x for x in other_parts if x is not None) or None
            if other is None and tot is not None and ath is not None:
                other = round(tot - ath, 2)
            put(rec, "athleticAid", ath)
            put(rec, "unfundedAid", other)

    if HS_NAME.exists():
        headers, rows = load_sheet(HS_NAME)
        i_id = col(headers, "PCID")
        i_hs = col(headers, "High School Name")
        for r in rows:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            rec = by_id.setdefault(pid, {"pcid": pid})
            put(rec, "highSchoolName", blank(r[i_hs] if i_hs is not None else ""))

    if INCOMING.exists():
        headers, rows = load_sheet(INCOMING)
        i_id = col(headers, "PowerCampus ID", "PCID")
        i_term = col(headers, "Application Term")
        i_gpa = col(headers, "Official AMU GPA")
        i_act = col(headers, "ACT Superscore (verified only)")
        i_sat = col(headers, "SAT Superscore (verified only)")
        i_clt = col(headers, "CLT Superscore (verified only)")
        i_faith = col(headers, "Person Interview - Faith (24/25)")
        i_courage = col(headers, "Person Interview - Courage (24/25)")
        i_desire = col(headers, "Person Interview - Desire (24/25)")
        i_hsinv = col(headers, "Person Interview - High School Involvement (24/25)")
        i_career = col(headers, "Person Interview - College and Career Expectations (24/25)")
        i_plans = col(headers, "Person Interview - Plans and potential while at AMU(24/25)")
        i_accom = col(headers, "Person Interview - AC Comments (24/25)")
        i_sector = col(headers, "Person Recruit - Sector")
        i_major = col(headers, "Application Major")
        i_type = col(headers, "School Type")
        i_hs = col(headers, "School Name")
        i_zip = col(headers, "Address by Rank Overall US 5-digit ZIP Code")
        i_city = col(headers, "Address by Rank Overall City")
        i_state = col(headers, "Person State", "Address 1 Region")
        scored = []
        for r in rows:
            pid = norm_id(r[i_id] if i_id is not None else "")
            if not pid:
                continue
            scored.append((term_rank(str(r[i_term] or "") if i_term is not None else ""), pid, r))
        scored.sort(reverse=True)
        seen = set()
        for _, pid, r in scored:
            if pid in seen:
                continue
            seen.add(pid)
            rec = by_id.setdefault(pid, {"pcid": pid})
            put(rec, "hsGpa", score(r[i_gpa] if i_gpa is not None else None))
            put(rec, "act", score(r[i_act] if i_act is not None else None), overwrite=True)
            put(rec, "sat", score(r[i_sat] if i_sat is not None else None, sat=True), overwrite=True)
            put(rec, "clt", score(r[i_clt] if i_clt is not None else None), overwrite=True)
            term = str(r[i_term] or "") if i_term is not None else ""
            is_fa25 = "2025" in term.lower() and "fall" in term.lower()
            # FA24/25 scores for the FA25 class only. Skip FA25/26 rubric and the
            # extract's Interview Total Rating (it mixes 25/26 for Fall 2026).
            if is_fa25:
                faith = score(r[i_faith] if i_faith is not None else None)
                courage = score(r[i_courage] if i_courage is not None else None)
                desire = score(r[i_desire] if i_desire is not None else None)
                hsinv = score(r[i_hsinv] if i_hsinv is not None else None)
                career = score(r[i_career] if i_career is not None else None)
                plans = score(r[i_plans] if i_plans is not None else None)
                put(rec, "interviewFaith", faith, overwrite=True)
                put(rec, "interviewCourage", courage, overwrite=True)
                put(rec, "interviewDesire", desire, overwrite=True)
                put(rec, "interviewHsInvolvement", hsinv, overwrite=True)
                put(rec, "interviewCollegeCareer", career, overwrite=True)
                put(rec, "interviewPlans", plans, overwrite=True)
                parts = [v for v in (faith, courage, desire, hsinv, career, plans) if v is not None]
                if parts:
                    total = sum(parts)
                    put(rec, "interviewTotal", int(total) if total == int(total) else total, overwrite=True)
                put(rec, "acComments", blank(r[i_accom] if i_accom is not None else ""), overwrite=True)
            put(rec, "sector", blank(r[i_sector] if i_sector is not None else ""), overwrite=True)
            put(rec, "applicationMajor", blank(r[i_major] if i_major is not None else ""), overwrite=True)
            put(rec, "schoolType", blank(r[i_type] if i_type is not None else ""), overwrite=True)
            put(rec, "highSchoolName", blank(r[i_hs] if i_hs is not None else ""), overwrite=True)
            put(rec, "zip", zip5(r[i_zip] if i_zip is not None else ""), overwrite=True)
            put(rec, "city", blank(r[i_city] if i_city is not None else ""), overwrite=True)
            put(rec, "state", blank(r[i_state] if i_state is not None else ""), overwrite=True)

    name_index: dict[tuple[str, str], dict] = {}
    if WITHDRAWAL.exists():
        headers, rows = load_sheet(WITHDRAWAL)
        i_fn = col(headers, "FirstName")
        i_ln = col(headers, "LastName")
        i_plan = col(headers, "FuturePlans")
        i_main = col(headers, "MainAreaOfImprovement")
        i_text = col(headers, "TextArea1")
        i_when = col(headers, "SubmitDate")
        i_year = col(headers, "Yearwithdrawing")
        scored_w = []
        for r in rows:
            fn = blank(r[i_fn] if i_fn is not None else "").lower()
            ln = blank(r[i_ln] if i_ln is not None else "").lower()
            if not fn or not ln:
                continue
            when = r[i_when] if i_when is not None else None
            ts = when.timestamp() if isinstance(when, datetime) else 0
            scored_w.append((ts, fn, ln, r))
        scored_w.sort(reverse=True)
        seen_n = set()
        for _, fn, ln, r in scored_w:
            key = (fn, ln)
            if key in seen_n:
                continue
            seen_n.add(key)
            name_index[key] = {
                "reasonForLeaving": blank(r[i_plan] if i_plan is not None else ""),
                "withdrawalImprovement": blank(r[i_main] if i_main is not None else ""),
                "withdrawalComments": blank(r[i_text] if i_text is not None else ""),
                "withdrawalYear": blank(r[i_year] if i_year is not None else ""),
            }

    zips, cities = ensure_geo()
    if PRIOR_JSON.exists():
        prior = json.loads(PRIOR_JSON.read_text())
        for s in prior.get("students") or prior:
            pid = norm_id(s.get("pcid"))
            if not pid:
                continue
            rec = by_id.setdefault(pid, {"pcid": pid})
            if not rec.get("city"):
                rec["city"] = blank(s.get("city"))
            if not rec.get("state"):
                rec["state"] = blank(s.get("state"))
    miles_n = 0
    for rec in by_id.values():
        miles = miles_from_campus(zips, cities, rec.get("zip") or "", rec.get("city") or "", rec.get("state") or "")
        if miles is None:
            continue
        rec["milesFromCampus"] = miles
        miles_n += 1

    OUT.write_text(
        json.dumps(
            {"generated": datetime.now().strftime("%Y-%m-%d"), "byPcid": by_id, "byName": {f"{a}|{b}": v for (a, b), v in name_index.items()}},
            indent=2,
            default=str,
        )
        + "\n"
    )
    print(json.dumps({"out": str(OUT), "pcids": len(by_id), "withdrawalNames": len(name_index), "miles": miles_n}))


if __name__ == "__main__":
    main()
