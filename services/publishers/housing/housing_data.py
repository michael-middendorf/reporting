"""Housing occupancy + residents payload (ported from housing-dashboard/serve.py).

Source of truth is the eRezLife assignment API (not Qlik housing_erez_life.csv).
Qlik/Slate still supply gender, deposits, and the optional roster.
"""
from __future__ import annotations

import asyncio
import csv
import datetime
import http.cookiejar
import io
import json
import os
import uuid
from pathlib import Path
from typing import Optional
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, Request, build_opener, urlopen

try:
    import websockets

    _WEBSOCKETS = True
except ImportError:
    _WEBSOCKETS = False

API_KEY = os.getenv("EREZLIFE_API_KEY") or os.getenv("API_KEY", "")
SITE_URL = os.getenv("SITE_URL", "https://avemaria.erezlife.com/")
TARGET_SESSION = os.getenv("TARGET_SESSION", "2026FALL01")
EREZ_EXTRA_SESSIONS = [
    s.strip()
    for s in os.getenv("EREZ_EXTRA_SESSIONS", "2026FALL03Ireland").split(",")
    if s.strip()
]
ALLOWED_SESSIONS = {TARGET_SESSION, *EREZ_EXTRA_SESSIONS}
SLATE_TOKEN = os.getenv("SLATE_TOKEN", "")
SLATE_QUERY_URL = os.getenv("SLATE_QUERY_URL", "")
SLATE_DEPOSITS_TOKEN = os.getenv("SLATE_DEPOSITS_TOKEN", "")
SLATE_DEPOSITS_URL = os.getenv("SLATE_DEPOSITS_URL", "")
SLATE_ROSTER_URL = os.getenv("SLATE_ROSTER_URL", "")
QLIK_API_KEY = os.getenv("QLIK_API_KEY", "")
QLIK_TENANT = os.getenv("QLIK_TENANT", "")
QLIK_SPACE = os.getenv("QLIK_SPACE", "Digital Strategy Team (DEV)")
DATA_DIR = Path(os.getenv("DATA_DIR", "/tmp"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DATA_BUCKET = (os.getenv("DATA_BUCKET") or "").strip()
HOUSING_CACHE_PREFIX = os.getenv("HOUSING_CACHE_PREFIX", "cache/housing").strip("/")
UPLOAD_PREFIX = f"{HOUSING_CACHE_PREFIX}/uploads"
RESIDENTS_KEY = f"{HOUSING_CACHE_PREFIX}/residents.json"
EREZ_CACHE_FILE = DATA_DIR / "erez_cache.json"
ROSTER_CACHE_FILE = DATA_DIR / "roster_cache.json"
DEPOSITS_CSV_FILE = DATA_DIR / "slate_deposits.csv"
CACHE_TTL_HOURS = float(os.getenv("HOUSING_CACHE_TTL_HOURS", "6"))

# eRez assignment fields. price_type_id is how Rome is stored: halls=UGDP,
# commuter=NODM, Ireland=ISTA (CIARA/MALAC), Rome=RSTA (no building_id),
# commuter study-abroad=CSTA.
_EREZ_FIELDS = [
    "session_id",
    "student_id",
    "email",
    "application_status",
    "start_date",
    "end_date",
    "building_id",
    "room_id",
    "price_type_id",
]
_ADMISSIONS_GENDER_FIELDS = ["Person Reference ID", "Gender"]
_DEPOSITS_FIELDS = [
    "Dep_App Created Date",
    "Dep_App Submitted Date",
    "Dep_Application Term",
    "Dep_Application Source",
    "Dep_Application Status",
    "Dep_Application Type",
    "Dep_Application Student Type",
    "Dep_Application Intended Major",
    "Dep_Application Intended Minor",
    "Dep_Application AC Assignment",
    "Dep_Application Workflow Bin",
    "Dep_Application Decision Code",
    "Dep_Application Decision Status",
    "Dep_Application Deposited Date",
    "Dep_Application Accepted Date",
    "Dep_Application Denied Date",
    "Dep_Application Withdrawn Date",
    "Dep_Application Athletic Decision",
    "Dep_Sport",
    "Dep_City",
    "Dep_Postal",
    "Dep_Region",
    "Dep_County",
    "Dep_Commuter Distance",
    "Dep_Visit Status",
    "Dep_First Name",
    "Dep_Last Name",
    "Dep_Self_SAT",
    "Dep_Self_ACT",
    "Dep_Self_CLT",
    "Portal Link",
    "Dep_GPA Self Reported",
    "Dep_Religion",
    "Dep_Official AMU GPA",
    "Dep_PCID",
    "Dep_Recruit - Rating",
    "Dep_Birthdate",
    "Dep_Phone",
    "Dep_LOI Status",
    "Dep_Offered",
    "Dep_CLT Convert  to ACT",
    "Dep_SAT Convert to ACT",
    "Dep_Application GUID",
    "Dep_LOI Released Date",
    "Dep_Deposit Date",
    "Dep_Withdrawn Date",
    "Dep_LOI Expire Date",
    "Dep_LOI Accepted Date",
    "Dep_Gender",
    "Dep_Risk Factor",
]
_ADMISSIONS_FIELDS = [
    "Application Student Type",
    "Application Type",
    "Application Status",
    "Application AC Assignment",
    "Created Date",
    "Application Submitted Date",
    "Application Term",
    "Gender",
    "Religion",
    "Sport",
    "Athlete or Non Athlete",
    "Application Athletic Decision",
    "Athlete Recruit - Rating",
    "SAT",
    "ACT",
    "CLT",
    "All Decision Codes",
    "Decision Code",
    "Accept Code",
    "Accept Date",
    "Full Accept Date",
    "Full Accept Code",
    "Deposit Code",
    "Deposit Date",
    "Denied Code",
    "Denied Date",
    "Withdrawn Code",
    "Withdrawn Date",
    "Self-Reported HS GPA",
    "Scholarship Amount",
    "Person Interview - Total Rating",
    "Address 1 City",
    "Address 1 Region",
    "Address 1 Postal",
    "Address 1 Country",
    "Person Reference ID",
    "Application Reference ID",
    "First Source Date",
    "First Source Memo",
    "Source Type Type",
    "Interview Status",
    "Interview Scheduled",
    "Interviewed Date FA25 Apps",
    "Interviewed Date FA26 Apps",
]

_INCOMING_TERMS = {"Fall 2026", "Summer 2026"}
_ATH_STATUS = {"Roster - Athlete", "Roster - Athletic Associated"}
_ROSTER_STATUS_OK = {
    "Roster - Athlete",
    "Roster - Athletic Associated",
    "Roster - Athlete Post Grad",
}
_MANAGER_TYPE = "Manager/Student Coach"
UPLOAD_KINDS = {"registration": "registration.xlsx", "early-arrivals": "early_arrivals.xlsx"}
_CLASS_MAP = {
    "FRES": "Freshman",
    "FIRS": "First Year",
    "SOPH": "Sophomore",
    "SECO": "Second Year",
    "JR": "Junior",
    "SR": "Senior",
    "GRAD": "Graduate",
    "NDEG": "Non-Degree",
}
_ATHLETIC_TABS = {
    "ATHLETIC EARLY ARRIVALS - 8.2",
    "FB - 8.3",
    "VB - 8.3",
    "BAND - 8.3",
    "MSOC - 8.3",
    "WSOC - 8.3",
    "XC - 8.3",
}
_EA_SKIP_TABS = {
    "Schedule",
    "MASTER",
    "Student Look Up",
    "Sum Res 7.20",
    "XAVIER",
    "GORET",
    "SEBAST",
    "STJOSE",
    "TERESA",
    "JPII",
}
_GROUP_LABELS = {
    "FB": "Football",
    "VB": "Volleyball",
    "BAND": "Band",
    "MSOC": "Men's Soccer",
    "WSOC": "Women's Soccer",
    "XC": "Cross Country",
    "ATHLETIC EARLY ARRIVALS": "Athletics",
}
DEFAULT_NEW, DEFAULT_RETURN = "2026-08-17", "2026-08-19"
_UPLOAD_KEYWORDS = {
    "registration.xlsx": ("regist",),
    "early_arrivals.xlsx": ("early", "arriv"),
}


def _normalize_dep(name: str) -> str:
    """'Dep_Application Decision Code' → 'Dep_ApplicationDecisionCode'"""
    if name.startswith("Dep_") and " " in name:
        return "Dep_" + name[4:].replace(" ", "")
    return name


def _display_date(dt: Optional[datetime.datetime] = None) -> str:
    dt = dt or datetime.datetime.now()
    return f"{dt.strftime('%b')} {dt.day} {dt.year}"


def _s3():
    import boto3

    region = os.getenv("AWS_REGION") or os.getenv("AWS_DEFAULT_REGION") or "us-east-2"
    return boto3.client("s3", region_name=region)


def _store_get(fname: str, *, prefix: str | None = None) -> Optional[bytes]:
    """Bytes of a stored upload/state object. S3 first, then DATA_DIR."""
    key_prefix = (prefix or UPLOAD_PREFIX).strip("/")
    if DATA_BUCKET:
        try:
            client = _s3()
            try:
                return client.get_object(Bucket=DATA_BUCKET, Key=f"{key_prefix}/{fname}")["Body"].read()
            except Exception:
                pass
            kw = _UPLOAD_KEYWORDS.get(fname, ())
            if kw:
                resp = client.list_objects_v2(Bucket=DATA_BUCKET, Prefix=f"{key_prefix}/")
                for obj in resp.get("Contents", []):
                    low = obj["Key"].lower()
                    if low.endswith(".xlsx") and any(w in low for w in kw):
                        return client.get_object(Bucket=DATA_BUCKET, Key=obj["Key"])["Body"].read()
        except Exception as exc:
            print(f"  S3 get {fname}: {exc}")
    path = DATA_DIR / fname
    return path.read_bytes() if path.exists() else None


def write_private(path: Path, text: str) -> None:
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w") as fh:
        fh.write(text)
    os.chmod(path, 0o600)


async def _qlik_fetch_async(filename: str, fields: list) -> list:
    session_id = str(uuid.uuid4())
    url = f"wss://{QLIK_TENANT}/app/SessionApp_{session_id}"
    headers = {"Authorization": f"Bearer {QLIK_API_KEY}"}
    mid = 0

    async def rpc(ws, method, handle, params):
        nonlocal mid
        mid += 1
        m = mid
        await ws.send(
            json.dumps({"jsonrpc": "2.0", "id": m, "method": method, "handle": handle, "params": params})
        )
        while True:
            raw = await asyncio.wait_for(ws.recv(), timeout=90)
            resp = json.loads(raw)
            if resp.get("id") == m:
                return resp

    async with websockets.connect(url, additional_headers=headers, open_timeout=20) as ws:
        await asyncio.wait_for(ws.recv(), timeout=15)

        resp = await rpc(ws, "OpenDoc", -1, {"qDocName": f"SessionApp_{session_id}"})
        dh = resp["result"]["qReturn"]["qHandle"]

        if filename.endswith(".csv"):
            fmt = "txt, codepage is 65001, embedded labels, delimiter is ',', msq"
        else:
            fmt = "qvd"
        script = f"data: LOAD * FROM [lib://{QLIK_SPACE}:DataFiles/{filename}] ({fmt});"
        await rpc(ws, "SetScript", dh, {"qScript": script})
        resp = await rpc(ws, "DoReload", dh, {})
        if not resp["result"].get("qReturn"):
            raise RuntimeError(f"Qlik reload failed for {filename}")

        dims = [{"qDef": {"qFieldDefs": [f], "qSortCriterias": []}} for f in fields]
        resp = await rpc(
            ws,
            "CreateObject",
            dh,
            {
                "qProp": {
                    "qInfo": {"qType": "table"},
                    "qHyperCubeDef": {
                        "qDimensions": dims,
                        "qMeasures": [],
                        "qInitialDataFetch": [{"qTop": 0, "qLeft": 0, "qHeight": 0, "qWidth": 0}],
                    },
                }
            },
        )
        oh = resp["result"]["qReturn"]["qHandle"]

        resp = await rpc(ws, "GetLayout", oh, {})
        total = resp["result"]["qLayout"]["qHyperCube"]["qSize"]["qcy"]
        print(f"    Qlik {filename}: {total} rows")

        records = []
        page = max(1, 10_000 // max(len(fields), 1))
        for top in range(0, total, page):
            h = min(page, total - top)
            resp = await rpc(
                ws,
                "GetHyperCubeData",
                oh,
                {
                    "qPath": "/qHyperCubeDef",
                    "qPages": [{"qTop": top, "qLeft": 0, "qHeight": h, "qWidth": len(fields)}],
                },
            )
            if "error" in resp:
                raise RuntimeError(f"GetHyperCubeData error: {resp['error']}")
            for row in resp["result"]["qDataPages"][0]["qMatrix"]:
                records.append({fields[i]: row[i].get("qText", "") for i in range(len(fields))})

        return records


def _qlik_fetch(filename: str, fields: list) -> list:
    if not _WEBSOCKETS:
        raise RuntimeError("websockets package not installed")
    if not QLIK_API_KEY or not QLIK_TENANT:
        raise RuntimeError("QLIK_API_KEY / QLIK_TENANT not set")
    return asyncio.run(_qlik_fetch_async(filename, fields))


def _erez_api_fetch(session_id: str) -> list[dict]:
    """Pull housing assignments for one session from the eRez API."""
    if not API_KEY:
        raise RuntimeError("EREZLIFE_API_KEY/API_KEY not set — cannot fetch from eRez API")
    base = SITE_URL if SITE_URL.endswith("/") else SITE_URL + "/"
    endpoint = base + "api/housing/assignment/"
    hdrs = {"x-api-key": API_KEY, "Accept": "application/json"}

    def _get(page=None, attempts=3):
        params = {"session_id": session_id}
        if page is not None:
            params["page"] = page
        last_exc = None
        for i in range(attempts):
            try:
                req = Request(f"{endpoint}?{urlencode(params)}", headers=hdrs)
                with urlopen(req, timeout=180) as resp:
                    return json.loads(resp.read().decode("utf-8"))
            except Exception as exc:
                last_exc = exc
                print(f"  … eRez {session_id} page {page or 1}: attempt {i + 1}/{attempts} failed ({exc})")
        raise last_exc

    data = _get()
    content = list(data.get("content", []))
    for page in [p["page"] for p in data.get("pages", [])][1:]:
        content.extend(_get(page).get("content", []))
    return [
        {k: rec.get(k, "") for k in _EREZ_FIELDS}
        for rec in content
        if rec.get("session_id") == session_id
    ]


def _gather_housing_records() -> list[dict]:
    records: list[dict] = []
    for sess in sorted(ALLOWED_SESSIONS):
        rows = _erez_api_fetch(sess)
        print(f"  → eRez API: {len(rows)} records for {sess}")
        records.extend(rows)
    return records


def _result_is_complete(result: dict) -> bool:
    return any(r.get("session_id") == TARGET_SESSION for r in result.get("records", []))


def _build_gender_lookup() -> dict:
    try:
        rows = _qlik_fetch("admissions_report_housing.csv", _ADMISSIONS_GENDER_FIELDS)
        lookup = {}
        for r in rows:
            pid = (r.get("Person Reference ID", "") or "").strip().upper()
            gen = (r.get("Gender", "") or "").strip()[:1].upper()
            if pid and gen:
                lookup[pid] = gen
        print(f"  → Gender lookup built: {len(lookup)} students")
        return lookup
    except Exception as exc:
        print(f"  ⚠  Gender lookup failed: {exc}")
        return {}


def _enrich_gender(records: list[dict], lookup: dict) -> list[dict]:
    for r in records:
        sid = (r.get("student_id", "") or "").strip().upper()
        r["gender"] = lookup.get(sid, "")
    return records


def _withdrawn_student_ids() -> set:
    raw = None
    try:
        raw = [{_normalize_dep(k): v for k, v in r.items()} for r in _qlik_fetch("housing_all_depoits.csv", _DEPOSITS_FIELDS)]
    except Exception as exc:
        print(f"  withdrawn lookup: Qlik failed ({exc}); trying Slate…")
        try:
            if not SLATE_DEPOSITS_URL:
                return set()
            opener = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
            req = Request(SLATE_DEPOSITS_URL, headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0"})
            with opener.open(req, timeout=60) as r:
                data = json.loads(r.read())
            raw = data.get("row") or data.get("records") or (data if isinstance(data, list) else [])
        except Exception as exc2:
            print(f"  withdrawn lookup: Slate also failed ({exc2}) — no exclusion applied")
            return set()
    out = set()
    for rec in raw or []:
        pcid = (rec.get("Dep_PCID") or "").strip().upper()
        if not pcid:
            continue
        code = (rec.get("Dep_ApplicationDecisionCode") or "").strip().upper()
        status = (rec.get("Dep_ApplicationStatus") or "").lower()
        wd = (rec.get("Dep_ApplicationWithdrawnDate") or rec.get("Dep_WithdrawnDate") or "").strip()
        if code == "WT" or wd or "withdrawn" in status:
            out.add(pcid)
    return out


def _build_housing_result(records: list[dict], include_gender: bool = False) -> dict:
    fa26 = [r for r in records if r.get("session_id") in ALLOWED_SESSIONS]
    withdrawn = _withdrawn_student_ids()
    if withdrawn:
        before = len(fa26)
        fa26 = [r for r in fa26 if (r.get("student_id") or "").strip().upper() not in withdrawn]
        dropped = before - len(fa26)
        if dropped:
            print(f"  → excluded {dropped} withdrawn student(s) from occupancy")
    if include_gender:
        _enrich_gender(fa26, _build_gender_lookup())
    return {
        "records": fa26,
        "session": TARGET_SESSION,
        "total": len(fa26),
        "cached_at": _display_date(),
        "source": "erez",
    }


def get_housing_data(*, force_refresh: bool = False) -> dict:
    """Cache-first housing occupancy payload from the eRez API."""
    if not force_refresh and EREZ_CACHE_FILE.exists():
        try:
            age_hours = (datetime.datetime.now().timestamp() - EREZ_CACHE_FILE.stat().st_mtime) / 3600
            if age_hours < CACHE_TTL_HOURS:
                return json.loads(EREZ_CACHE_FILE.read_text())
        except Exception as exc:
            print(f"  Cache read error: {exc}")

    records = _gather_housing_records()
    result = _build_housing_result(records, include_gender=True)
    if not _result_is_complete(result):
        raise RuntimeError(f"eRez fetch missing main session {TARGET_SESSION} — refusing to cache partial data")
    write_private(EREZ_CACHE_FILE, json.dumps(result))
    return result


def _name_from_email(email: str) -> str:
    local = (email or "").split("@")[0]
    parts = [p for p in local.replace("_", ".").split(".") if p]
    return " ".join(w.capitalize() for w in parts)


def _fetch_slate_roster() -> dict:
    if not SLATE_ROSTER_URL:
        return {}
    rows = None
    if ROSTER_CACHE_FILE.exists():
        try:
            age = (datetime.datetime.now().timestamp() - ROSTER_CACHE_FILE.stat().st_mtime) / 3600
            if age < 12:
                rows = json.loads(ROSTER_CACHE_FILE.read_text())
        except Exception:
            rows = None
    for attempt in range(2):
        if rows:
            break
        try:
            opener = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
            req = Request(SLATE_ROSTER_URL, headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0"})
            with opener.open(req, timeout=200) as r:
                rows = json.loads(r.read()).get("row", [])
            if rows:
                try:
                    write_private(ROSTER_CACHE_FILE, json.dumps(rows))
                except Exception:
                    pass
                break
        except Exception as exc:
            print(f"  roster fetch attempt {attempt + 1}/2 failed: {exc}")
    if not rows and ROSTER_CACHE_FILE.exists():
        try:
            rows = json.loads(ROSTER_CACHE_FILE.read_text())
            print("  roster: using last-good cache")
        except Exception:
            rows = None
    idx = {}
    for row in rows or []:
        p = (row.get("PCID") or "").strip().upper()
        e = (row.get("Email") or "").strip().lower()
        if p:
            idx[p] = row
        if e:
            idx["e:" + e] = row
    return idx


def _load_local_deposits_csv():
    if not DEPOSITS_CSV_FILE.exists():
        return None
    text = DEPOSITS_CSV_FILE.read_text(encoding="utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    records = list(reader)
    print(f"  Loaded {len(records)} from {DEPOSITS_CSV_FILE.name} (fallback)\n")
    return records


def _deposit_sets():
    raw = None
    try:
        raw = [{_normalize_dep(k): v for k, v in r.items()} for r in _qlik_fetch("housing_all_depoits.csv", _DEPOSITS_FIELDS)]
    except Exception:
        try:
            records = _load_local_deposits_csv()
            if records is not None:
                raw = records
            elif not SLATE_DEPOSITS_URL:
                return set(), set(), []
            else:
                opener = build_opener(HTTPCookieProcessor(http.cookiejar.CookieJar()))
                req = Request(SLATE_DEPOSITS_URL, headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0"})
                with opener.open(req, timeout=60) as r:
                    raw = json.loads(r.read()).get("row", [])
        except Exception:
            return set(), set(), []
    inc, wd = set(), set()
    for rec in raw or []:
        p = (rec.get("Dep_PCID") or "").strip().upper()
        if not p:
            continue
        code = (rec.get("Dep_ApplicationDecisionCode") or "").strip().upper()
        status = (rec.get("Dep_ApplicationStatus") or "").lower()
        wdt = (rec.get("Dep_ApplicationWithdrawnDate") or rec.get("Dep_WithdrawnDate") or "").strip()
        (wd if (code == "WT" or wdt or "withdrawn" in status) else inc).add(p)
    return inc, wd, (raw or [])


def _group_label(tab: str) -> str:
    base = tab.split(" - ")[0].strip()
    return _GROUP_LABELS.get(base.upper(), base)


def _sport_key(name: str) -> str:
    n = (name or "").strip().lower()
    for g in ("men's ", "women's ", "mens ", "womens "):
        n = n.replace(g, "")
    if "cross country" in n:
        return "cross country"
    if "band" in n:
        return "band"
    if "soccer" in n:
        return "soccer"
    return n.strip()


def _early_reason(groups: list[str], sport: str, is_ath: bool) -> str:
    parts: list[str] = []
    sport = (sport or "").strip()
    ath = [g for g in groups if g in _ATHLETIC_TABS]
    if ath:
        specific, has_generic = [], False
        for g in ath:
            if g.split(" - ")[0].strip().upper() == "ATHLETIC EARLY ARRIVALS":
                has_generic = True
            else:
                lbl = _group_label(g)
                if lbl not in specific:
                    specific.append(lbl)
        if specific:
            parts.extend(specific)
        elif is_ath and sport:
            parts.append(sport)
        elif has_generic:
            parts.append("Athletics")
    for g in groups:
        if g in _ATHLETIC_TABS:
            continue
        lbl = _group_label(g)
        if lbl not in parts:
            parts.append(lbl)
    return " · ".join(parts)


def _as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return None


def _parse_registration():
    data = _store_get("registration.xlsx")
    if not data:
        return {}, {}
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb["Registration Status"] if "Registration Status" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and "ID" in [str(c or "").strip() for c in r]), 0)
    hdr = [str(c or "").strip() for c in rows[hi]]

    def col(n):
        return hdr.index(n) if n in hdr else None

    idi, ci, ei, pf, coi = col("ID"), col("Class"), col("Email"), col("Primary Flag"), col("College Attend")
    by_p, by_e = {}, {}
    for r in rows[hi + 1 :]:
        if not r or idi is None or idi >= len(r) or not r[idi]:
            continue
        pcid = str(r[idi]).strip().upper()
        cls = str(r[ci] or "").strip() if ci is not None and ci < len(r) else ""
        coll = str(r[coi] or "").strip().upper() if coi is not None and coi < len(r) else ""
        v = {"cls": _CLASS_MAP.get(cls, cls.title()) if cls else "", "coll": coll}
        prim = (str(r[pf] or "").strip().upper() == "Y") if pf is not None and pf < len(r) else False
        if pcid not in by_p or prim:
            by_p[pcid] = v
        em = str(r[ei] or "").strip().lower() if ei is not None and ei < len(r) else ""
        if em and (em not in by_e or prim):
            by_e[em] = v
    return by_p, by_e


def _parse_early_arrivals():
    data = _store_get("early_arrivals.xlsx")
    if not data:
        return {}
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    npc = lambda x: str(x or "").strip().upper()
    look = {}
    if "Student Look Up" in wb.sheetnames:
        for r in list(wb["Student Look Up"].iter_rows(values_only=True))[1:]:
            if r and r[0]:
                look[npc(r[0])] = {
                    "first": r[1],
                    "last": r[2],
                    "email": r[3] if len(r) > 3 else "",
                    "category": r[6] if len(r) > 6 else "",
                    "class": r[7] if len(r) > 7 else "",
                }
    out = {}
    for t in wb.sheetnames:
        if t in _EA_SKIP_TABS:
            continue
        rows = list(wb[t].iter_rows(values_only=True))
        if not rows:
            continue
        hdr = [str(c or "").strip().upper() for c in rows[0]]
        try:
            pci = hdr.index("PCID#")
        except ValueError:
            continue
        di = next((i for i, h in enumerate(hdr) if "MOVE-IN" in h), 8)
        for r in rows[1:]:
            if not r or pci >= len(r) or not r[pci]:
                continue
            p = npc(r[pci])
            info = out.setdefault(p, {"groups": [], "dates": {}, "look": look.get(p, {})})
            info["groups"].append(t)
            d = _as_date(r[di]) if di < len(r) else None
            if d:
                info["dates"][t] = d
    return out


def _build_cloud_residents(erez_records: list[dict]) -> dict:
    npc = lambda x: str(x or "").strip().upper()
    roster = _fetch_slate_roster()
    roster_loaded = bool(roster)
    incoming, withdrawn, dep_rows = _deposit_sets()
    try:
        gender_lookup = _build_gender_lookup()
    except Exception:
        gender_lookup = {}
    reg_p, reg_e = _parse_registration()
    ea = _parse_early_arrivals()
    residents, candidates, excluded = [], [], 0
    for rec in erez_records:
        pcid = npc(rec.get("student_id"))
        if pcid in withdrawn:
            excluded += 1
            continue
        email = rec.get("email", "")
        s = roster.get(pcid) or roster.get("e:" + email.strip().lower()) or {}
        et = (s.get("EntryTerm") or "").strip()
        st = (s.get("Status") or "").strip().lower()
        is_ath = bool(s.get("RosterStatus") in _ATH_STATUS)
        regv = reg_p.get(pcid) or reg_e.get(email.strip().lower()) or {}
        cls = regv.get("cls", "")
        if cls:
            fresh = cls.lower().startswith("fresh") or cls.lower() in ("first year", "new")
            et_inc = et in _INCOMING_TERMS
            inc = fresh or et_inc
            transfer = (regv.get("coll") == "TRAN") and not fresh
            stype = "Incoming" if fresh else ("Incoming Transfer" if et_inc else "Returning")
        else:
            inc = pcid in incoming or st == "applicant" or et in _INCOMING_TERMS
            transfer = False
            stype = "Incoming" if inc else "Returning"
        first = (s.get("First") or "").strip()
        last = (s.get("Last") or "").strip()
        info = ea.get(pcid)
        look = (info.get("look") or {}) if info else {}
        if not (first or last) and look:
            first = str(look.get("first") or "").strip()
            last = str(look.get("last") or "").strip()
        name = (first + " " + last).strip() or _name_from_email(email) or pcid
        gender = (s.get("Gender") or "").strip()[:1].upper() or gender_lookup.get(pcid, "")
        move_in = DEFAULT_NEW if inc else DEFAULT_RETURN
        move_in_source, move_in_early, early_reason, early_groups = "class-default", "", "", []
        if info:
            applicable = [(d, t) for t, d in info["dates"].items()]
            if applicable:
                best = min(applicable)
                move_in, move_in_source, move_in_early = best[0], best[1], best[0]
            groups = sorted(set(info["groups"]))
            early_groups = groups
            early_reason = _early_reason(early_groups, s.get("Sport") or "", is_ath)
            sheet_sports, slate_sport = [], (s.get("Sport") or "").strip()
            for g in groups:
                if g in _ATHLETIC_TABS and g.split(" - ")[0].strip().upper() != "ATHLETIC EARLY ARRIVALS":
                    lbl = _group_label(g)
                    if lbl not in sheet_sports:
                        sheet_sports.append(lbl)
            on_ath = any(g in _ATHLETIC_TABS for g in groups)
            roster_status = (s.get("RosterStatus") or "").strip()
            ath_type = (s.get("AthleticType") or "").strip()
            reasons = []
            if (
                roster_loaded
                and slate_sport
                and sheet_sports
                and not any(_sport_key(ss) == _sport_key(slate_sport) for ss in sheet_sports)
            ):
                reasons.append(
                    f"Sport mismatch — early-arrival sheet lists {' / '.join(sheet_sports)}, "
                    f"Slate roster has {slate_sport}"
                )
            if roster_loaded and on_ath and roster_status not in _ROSTER_STATUS_OK and ath_type != _MANAGER_TYPE:
                if not s:
                    detail = "not found in Slate"
                elif roster_status:
                    detail = f"Slate status: {roster_status}"
                else:
                    detail = "no roster standing in Slate"
                reasons.append(f"Not on the athletic roster or a student manager — {detail}")
            if reasons:
                dd = info["dates"]
                earliest = min(dd.items(), key=lambda kv: kv[1])[0] if dd else (groups[0] if groups else "")
                candidates.append(
                    {
                        "pcid": pcid,
                        "name": name,
                        "email": email,
                        "building": rec.get("building_id", ""),
                        "room": rec.get("room_id", ""),
                        "category": look.get("category", "") or "",
                        "class_status": cls,
                        "student_type": stype,
                        "gender": gender,
                        "roster_status": roster_status,
                        "athletic_type": ath_type,
                        "is_athlete": is_ath,
                        "sheet_sport": " / ".join(sheet_sports),
                        "slate_sport": slate_sport,
                        "groups": [earliest],
                        "move_in": move_in,
                        "move_in_source": move_in_source,
                        "reasons": reasons,
                    }
                )
        residents.append(
            {
                "pcid": pcid,
                "email": email,
                "first": first,
                "last": last,
                "name": name,
                "building": rec.get("building_id", ""),
                "room": rec.get("room_id", ""),
                "session": rec.get("session_id", ""),
                "class_status": cls,
                "student_type": stype,
                "transfer": transfer,
                "college_attend": regv.get("coll", ""),
                "category": look.get("category", "") or "",
                "sport": (s.get("Sport") or ""),
                "roster_status": (s.get("RosterStatus") or ""),
                "athletic_type": (s.get("AthleticType") or ""),
                "entry_term": et,
                "gender": gender,
                "hometown": ", ".join(x for x in [(s.get("City") or ""), (s.get("State") or "")] if x),
                "is_athlete": is_ath,
                "price_type": rec.get("price_type_id", ""),
                "move_in_erez": rec.get("start_date") or "",
                "move_in_early": move_in_early,
                "move_in": move_in,
                "move_in_source": move_in_source,
                "end_date": rec.get("end_date") or "",
                "early_reason": early_reason,
                "early_groups": early_groups,
            }
        )
    candidates.sort(key=lambda c: (c["name"] or "").lower())
    g = lambda r: (r.get("Dep_Gender") or "").strip()[:1].upper()
    wt = [r for r in dep_rows if r.get("Dep_ApplicationDecisionCode") == "WT"]
    return {
        "generated": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "count": len(residents),
        "source": "erez",
        "session": TARGET_SESSION,
        "cached_at": _display_date(),
        "withdrawn": {
            "excluded_from_feed": excluded,
            "deposit_total": len(dep_rows),
            "deposit_total_w": sum(1 for r in dep_rows if g(r) == "F"),
            "deposit_total_m": sum(1 for r in dep_rows if g(r) == "M"),
            "withdrawn_total": len(wt),
            "withdrawn_w": sum(1 for r in wt if g(r) == "F"),
            "withdrawn_m": sum(1 for r in wt if g(r) == "M"),
        },
        "verification_candidates": candidates,
        "residents": residents,
    }


def build_residents_payload(occupancy: dict | None = None) -> dict:
    erez = (occupancy or {}).get("records")
    if not erez:
        erez = (
            json.loads(EREZ_CACHE_FILE.read_text()).get("records", [])
            if EREZ_CACHE_FILE.exists()
            else _gather_housing_records()
        )
    return _build_cloud_residents(erez)


def study_abroad_from_records(records: list[dict]) -> dict:
    houses: dict[str, int] = {}
    for r in records:
        b = (r.get("building_id") or r.get("building") or "").strip()
        if b in ("CIARA", "MALAC"):
            houses[b] = houses.get(b, 0) + 1
    return {
        "session": ",".join(sorted(EREZ_EXTRA_SESSIONS)),
        "houses": houses,
        "total": sum(houses.values()),
        "cached_at": _display_date(),
    }


def get_admissions_for_housing() -> dict:
    try:
        records = _qlik_fetch("admissions_report_housing.csv", _ADMISSIONS_FIELDS)
        return {"records": records, "total": len(records)}
    except Exception as exc:
        if not SLATE_QUERY_URL:
            raise
        print(f"  Qlik admissions error: {exc}; trying Slate")
        jar = http.cookiejar.CookieJar()
        opener = build_opener(HTTPCookieProcessor(jar))
        req = Request(SLATE_QUERY_URL, headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0"})
        with opener.open(req, timeout=180) as r:
            data = json.loads(r.read())
        records = data.get("row") or data.get("records") or (data if isinstance(data, list) else [])
        return {"records": records, "total": len(records)}


def get_deposits_for_housing() -> dict:
    try:
        raw = _qlik_fetch("housing_all_depoits.csv", _DEPOSITS_FIELDS)
        records = [{_normalize_dep(k): v for k, v in r.items()} for r in raw]
        return {"records": records, "total": len(records), "source": "qlik"}
    except Exception as exc:
        print(f"  Qlik deposits error: {exc}")
        records = _load_local_deposits_csv()
        if records is not None:
            return {"records": records, "total": len(records), "source": "csv"}
        if not SLATE_DEPOSITS_URL:
            raise
        jar = http.cookiejar.CookieJar()
        opener = build_opener(HTTPCookieProcessor(jar))
        req = Request(SLATE_DEPOSITS_URL, headers={"Accept": "application/json", "User-Agent": "Mozilla/5.0"})
        with opener.open(req, timeout=60) as r:
            data = json.loads(r.read())
        records = data.get("row") or data.get("records") or (data if isinstance(data, list) else [])
        return {"records": records, "total": len(records), "source": "slate"}
